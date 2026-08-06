from __future__ import annotations

import json
import math
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from PIL import Image, ImageDraw
from scipy import stats

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
THERMAL_ROOT = ROOT.parent / "SeedThermal"
OUT = ROOT / "deliverables"
ASSETS = OUT / "_report_assets"
OUT.mkdir(exist_ok=True)
ASSETS.mkdir(exist_ok=True)

PLANT_CSV = ROOT / "results/combined_quality_mask_analysis_v1/plant_phenotypes.csv"
VARIETY_CSV = ROOT / "results/combined_quality_mask_analysis_v1/variety_phenotypes.csv"
STOMATA_CSV = ROOT / "results/quality_mask_analysis_v1/stomatal_complex_measurements.csv"
AREA_JSON = ROOT / "results/quality_mask_analysis_v1/area_comparisons/stomatal_area_analysis.json"
THERMAL_CSV = THERMAL_ROOT / "processed/thermal_data_all_sessions.csv"
UK_LOGO = THERMAL_ROOT / "dashboard/src/assets/images/uk_logo.png"

DOCX_PATH = OUT / "Integrated_Tomato_Phenotyping_Report.docx"

NAVY = "17324D"
INK = "1D2833"
TEAL = "1D8A8A"
BLUE = "2E7DA5"
GOLD = "C88A2B"
RED = "B64B4B"
MUTED = "60707E"
LIGHT = "F3F6F7"
PALE_BLUE = "EAF2F6"
PALE_TEAL = "E7F3F1"
PALE_GOLD = "FBF2E3"
WHITE = "FFFFFF"
GRID = "D6DEE3"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run(run, size=10.5, bold=False, italic=False, color=INK, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return run


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color=GRID, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run(run, 8.5, color=MUTED)


def paragraph(doc, text="", size=10.5, bold=False, italic=False, color=INK,
              align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, line=1.25,
              keep_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep_next
    if text:
        set_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def rich_paragraph(doc, parts, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT,
                   before=0, after=6, line=1.25, keep_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep_next
    for text, opts in parts:
        set_run(
            p.add_run(text),
            size=opts.get("size", size),
            bold=opts.get("bold", False),
            italic=opts.get("italic", False),
            color=opts.get("color", INK),
        )
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def kicker(doc, text):
    p = paragraph(doc, text.upper(), 8.5, True, color=TEAL, after=5, keep_next=True)
    for r in p.runs:
        r.font.all_caps = True
    return p


def source_note(doc, text):
    return paragraph(doc, text, 8, italic=True, color=MUTED, before=2, after=7, line=1.05)


def caption(doc, number: int, text: str):
    return rich_paragraph(
        doc,
        [(f"Figure {number}. ", {"bold": True, "color": NAVY}), (text, {"color": MUTED})],
        size=8.5,
        after=3,
        line=1.05,
        keep_next=True,
    )


def add_picture(doc, path: Path, width=6.1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(str(path), width=Inches(width))
    return p


def add_callout(doc, label: str, text: str, fill=PALE_TEAL, accent=TEAL):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color=fill, size=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.18
    set_run(p.add_run(label.upper() + "  "), 8.5, True, color=accent)
    set_run(p.add_run(text), 10.2, color=INK)
    paragraph(doc, "", after=4)


def add_metric_strip(doc, metrics):
    table = doc.add_table(rows=1, cols=len(metrics))
    widths = [9360 // len(metrics)] * len(metrics)
    widths[-1] += 9360 - sum(widths)
    set_table_geometry(table, widths)
    set_table_borders(table, color=WHITE, size=1)
    for cell, (value, label, fill, color) in zip(table.rows[0].cells, metrics):
        set_cell_shading(cell, fill)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        set_run(p.add_run(value), 17, True, color=color)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.0
        set_run(p2.add_run(label), 7.8, True, color=MUTED)
    paragraph(doc, "", after=4)


def add_key_value_table(doc, rows, widths=(2200, 7160)):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_shading(cells[0], PALE_BLUE)
        set_run(cells[0].paragraphs[0].add_run(label), 9, True, color=NAVY)
        set_run(cells[1].paragraphs[0].add_run(value), 9, color=INK)
        for c in cells:
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            c.paragraphs[0].paragraph_format.line_spacing = 1.1
    set_table_geometry(table, list(widths))
    set_table_borders(table)
    paragraph(doc, "", after=4)
    return table


def add_section_break(doc):
    doc.add_page_break()


def style_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    settings = {
        "Title": (30, NAVY, 0, 8),
        "Subtitle": (14, BLUE, 0, 10),
        "Heading 1": (16, NAVY, 18, 10),
        "Heading 2": (13, TEAL, 12, 6),
        "Heading 3": (11.5, NAVY, 8, 4),
    }
    for name, (size, color, before, after) in settings.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption_style = doc.styles["Caption"]
    caption_style.font.name = "Calibri"
    caption_style.font.size = Pt(8.5)
    caption_style.font.italic = False
    caption_style.font.color.rgb = rgb(MUTED)

    for sec in doc.sections:
        header = sec.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.space_after = Pt(0)
        set_run(hp.add_run("SAPKOTA LAB  |  INTEGRATED PHENOTYPING"), 7.5, True, color=MUTED)
        footer = sec.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.paragraph_format.space_after = Pt(0)
        set_run(fp.add_run("CONFIDENTIAL RESEARCH UPDATE  |  "), 7.5, True, color=MUTED)
        add_page_field(fp)


def configure_pyplot():
    plt.rcParams.update({
        "font.family": "DejaVu Sans",
        "font.size": 9,
        "axes.titlesize": 14,
        "axes.titleweight": "bold",
        "axes.labelsize": 9.5,
        "axes.edgecolor": "#60707E",
        "axes.linewidth": 0.8,
        "xtick.color": "#60707E",
        "ytick.color": "#60707E",
        "text.color": "#1D2833",
        "figure.facecolor": "white",
        "axes.facecolor": "white",
        "savefig.facecolor": "white",
    })


def generate_figures(plant: pd.DataFrame, variety: pd.DataFrame, stomata: pd.DataFrame):
    configure_pyplot()

    # Thermal paired comparison.
    t = variety.sort_values("difference_drought_minus_watered_canopy_mean_c")
    y = np.arange(len(t))
    fig, ax = plt.subplots(figsize=(8.2, 5.8))
    valid = t["watered_thermal_mask_qc_passed"] & t["drought_thermal_mask_qc_passed"]
    for i, (_, row) in enumerate(t.iterrows()):
        line_color = "#B9C3CA" if valid.iloc[i] else "#D7A6A6"
        ax.plot(
            [row["watered_canopy_mean_c"], row["drought_canopy_mean_c"]],
            [i, i], color=line_color, lw=1.6, zorder=1,
        )
    ax.scatter(t["watered_canopy_mean_c"], y, s=42, color="#" + BLUE, label="Watered", zorder=3)
    ax.scatter(t["drought_canopy_mean_c"], y, s=42, color="#" + GOLD, label="Drought-stressed", zorder=3)
    for i in np.where(~valid.to_numpy())[0]:
        ax.text(
            max(t.iloc[i]["watered_canopy_mean_c"], t.iloc[i]["drought_canopy_mean_c"]) + 0.08,
            i, "MASK QC", color="#" + RED, va="center", fontsize=7, fontweight="bold",
        )
    ax.set_yticks(y, t["variety"])
    ax.set_xlabel("Canopy mean temperature (°C)")
    ax.set_title("Canopy temperature varies by variety more than by treatment")
    ax.legend(frameon=False, ncol=2, loc="lower right")
    ax.grid(axis="x", color="#E9EEF1", lw=.8)
    ax.spines[["top", "right", "left"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(ASSETS / "thermal_variety_pairs.png", dpi=220, bbox_inches="tight")
    plt.close(fig)

    # Projected plant area treatment difference in valid pairs.
    q = t[valid].sort_values("difference_drought_minus_watered_projected_plant_area_px2")
    values = q["difference_drought_minus_watered_projected_plant_area_px2"].to_numpy()
    colors = ["#" + BLUE if x < 0 else "#" + GOLD for x in values]
    fig, ax = plt.subplots(figsize=(8.2, 5.2))
    yy = np.arange(len(q))
    ax.barh(yy, values, color=colors, height=.62)
    ax.axvline(0, color="#60707E", lw=1)
    ax.set_yticks(yy, q["variety"])
    ax.set_xlabel("Drought-stressed minus watered projected area (pixels²)")
    positive_pairs = int((values > 0).sum())
    ax.set_title(
        f"Projected plant area was larger under drought in {positive_pairs} of {len(values)} QC-valid pairs"
    )
    ax.grid(axis="x", color="#E9EEF1", lw=.8)
    ax.spines[["top", "right", "left"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(ASSETS / "projected_area_difference.png", dpi=220, bbox_inches="tight")
    plt.close(fig)

    # Stomatal treatment pairs.
    s = variety.sort_values("difference_drought_minus_watered")
    y = np.arange(len(s))
    fig, ax = plt.subplots(figsize=(8.2, 5.8))
    for i, (_, row) in enumerate(s.iterrows()):
        ax.plot(
            [row["watered_mean_area"], row["drought_mean_area"]],
            [i, i], color="#B9C3CA", lw=1.6, zorder=1,
        )
    ax.scatter(s["watered_mean_area"], y, s=42, color="#" + BLUE, label="Watered", zorder=3)
    ax.scatter(s["drought_mean_area"], y, s=42, color="#" + GOLD, label="Drought-stressed", zorder=3)
    ax.set_yticks(y, s["variety"])
    ax.set_xlabel("Plant mean stomatal-complex area (pixels²)")
    ax.set_title("Stomatal-complex area responses were heterogeneous across varieties")
    ax.legend(frameon=False, ncol=2, loc="lower right")
    ax.grid(axis="x", color="#E9EEF1", lw=.8)
    ax.spines[["top", "right", "left"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(ASSETS / "stomatal_variety_pairs.png", dpi=220, bbox_inches="tight")
    plt.close(fig)

    # Cross-modal relationships.
    qplant = plant[plant["thermal_mask_qc_passed"]].copy()
    fig, axes = plt.subplots(1, 2, figsize=(9.6, 4.0))
    for ax, x, xlabel in [
        (axes[0], "canopy_mean_c", "Canopy mean temperature (°C)"),
        (axes[1], "projected_plant_area_px2", "Projected plant area (pixels²)"),
    ]:
        for treatment, color, label in [
            ("watered", "#" + BLUE, "Watered"),
            ("drought_stressed", "#" + GOLD, "Drought-stressed"),
        ]:
            g = qplant[qplant["water_status"] == treatment]
            ax.scatter(g[x], g["stomatal_complex_area_mean"], s=46, color=color, alpha=.9, label=label)
        r, p = stats.pearsonr(qplant[x], qplant["stomatal_complex_area_mean"])
        ax.text(.04, .94, f"Pearson r = {r:.02f}\\np = {p:.3f}", transform=ax.transAxes,
                va="top", fontsize=8.5, color="#" + MUTED)
        ax.set_xlabel(xlabel)
        ax.set_ylabel("Mean stomatal-complex area (pixels²)")
        ax.grid(color="#E9EEF1", lw=.8)
        ax.spines[["top", "right"]].set_visible(False)
    axes[1].legend(frameon=False, loc="lower right", fontsize=8)
    fig.suptitle("No cross-sectional coupling between stomatal area and canopy-scale traits", fontsize=13, fontweight="bold")
    fig.tight_layout()
    fig.savefig(ASSETS / "cross_modal_scatter.png", dpi=220, bbox_inches="tight")
    plt.close(fig)

    # Manual mask example.
    row = stomata.sort_values("local_quality_score", ascending=False).iloc[0]
    image_path = ROOT / "Stomata_Images" / row["stomata_image"]
    image = Image.open(image_path).convert("RGB")
    polygon = json.loads(row["polygon_json"])
    xs = [p[0] for p in polygon]
    ys = [p[1] for p in polygon]
    pad = 90
    box = (
        max(0, int(min(xs) - pad)),
        max(0, int(min(ys) - pad)),
        min(image.width, int(max(xs) + pad)),
        min(image.height, int(max(ys) + pad)),
    )
    crop = image.crop(box)
    shifted = [(x - box[0], y - box[1]) for x, y in polygon]
    draw = ImageDraw.Draw(crop, "RGBA")
    draw.polygon(shifted, fill=(29, 138, 138, 55), outline=(255, 255, 255, 255), width=5)
    crop.save(ASSETS / "manual_mask_example.png", quality=95)


def compute_results(plant: pd.DataFrame, variety: pd.DataFrame):
    area_stats = json.loads(AREA_JSON.read_text())["paired_treatment_comparison"]

    thermal_all_diff = (
        variety["drought_canopy_mean_c"] - variety["watered_canopy_mean_c"]
    )
    thermal_all = {
        "drought_mean": variety["drought_canopy_mean_c"].mean(),
        "watered_mean": variety["watered_canopy_mean_c"].mean(),
        "difference": thermal_all_diff.mean(),
        "ci": stats.t.interval(
            .95, len(thermal_all_diff) - 1,
            loc=thermal_all_diff.mean(), scale=stats.sem(thermal_all_diff),
        ),
        "p": stats.ttest_rel(
            variety["drought_canopy_mean_c"], variety["watered_canopy_mean_c"]
        ).pvalue,
    }

    valid = variety[
        variety["watered_thermal_mask_qc_passed"]
        & variety["drought_thermal_mask_qc_passed"]
    ].copy()
    thermal_valid_diff = valid["difference_drought_minus_watered_canopy_mean_c"]
    thermal_valid = {
        "n": len(valid),
        "drought_mean": valid["drought_canopy_mean_c"].mean(),
        "watered_mean": valid["watered_canopy_mean_c"].mean(),
        "difference": thermal_valid_diff.mean(),
        "ci": stats.t.interval(
            .95, len(valid) - 1,
            loc=thermal_valid_diff.mean(), scale=stats.sem(thermal_valid_diff),
        ),
        "p": stats.ttest_rel(
            valid["drought_canopy_mean_c"], valid["watered_canopy_mean_c"]
        ).pvalue,
        "wilcoxon": stats.wilcoxon(thermal_valid_diff).pvalue,
    }
    plant_diff = valid["difference_drought_minus_watered_projected_plant_area_px2"]
    plant_area = {
        "n": len(valid),
        "drought_mean": valid["drought_projected_plant_area_px2"].mean(),
        "watered_mean": valid["watered_projected_plant_area_px2"].mean(),
        "difference": plant_diff.mean(),
        "pct": plant_diff.mean() / valid["watered_projected_plant_area_px2"].mean() * 100,
        "ci": stats.t.interval(
            .95, len(valid) - 1, loc=plant_diff.mean(), scale=stats.sem(plant_diff),
        ),
        "p": stats.ttest_rel(
            valid["drought_projected_plant_area_px2"],
            valid["watered_projected_plant_area_px2"],
        ).pvalue,
        "wilcoxon": stats.wilcoxon(plant_diff).pvalue,
        "positive_pairs": int((plant_diff > 0).sum()),
    }
    qplant = plant[plant["thermal_mask_qc_passed"]]
    cross = {
        "n": len(qplant),
        "temp": stats.pearsonr(
            qplant["canopy_mean_c"], qplant["stomatal_complex_area_mean"]
        ),
        "projected": stats.pearsonr(
            qplant["projected_plant_area_px2"], qplant["stomatal_complex_area_mean"]
        ),
        "response_temp": stats.pearsonr(
            valid["difference_drought_minus_watered_canopy_mean_c"],
            valid["difference_drought_minus_watered"],
        ),
        "response_area": stats.pearsonr(
            valid["difference_drought_minus_watered_projected_plant_area_px2"],
            valid["difference_drought_minus_watered"],
        ),
    }
    return area_stats, thermal_all, thermal_valid, plant_area, cross, valid


def add_cover(doc):
    p = paragraph(doc, "", after=0)
    p.paragraph_format.space_before = Pt(42)
    if UK_LOGO.exists():
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(UK_LOGO), width=Inches(1.0))
    paragraph(doc, "INTEGRATED TOMATO DROUGHT PHENOTYPING", 9, True, color=TEAL,
              align=WD_ALIGN_PARAGRAPH.CENTER, before=24, after=14)
    paragraph(doc, "From Canopy Heat\nto Stomatal Form", 30, True, color=NAVY,
              align=WD_ALIGN_PARAGRAPH.CENTER, after=12, line=1.0)
    paragraph(doc, "A combined SeedThermal and StomatalPhenotypingPipeline research report",
              14, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=26, line=1.15)
    paragraph(doc, "16 tomato varieties  |  32 plants  |  2 water treatments",
              10.5, True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, after=56)
    paragraph(doc, "Prepared for Sapkota Lab", 11, True, color=NAVY,
              align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
    paragraph(doc, "29 July 2026", 9.5, color=MUTED,
              align=WD_ALIGN_PARAGRAPH.CENTER, after=16)
    add_callout(
        doc,
        "Scope",
        "An integrated, quality-aware interpretation of thermal canopy traits, projected plant area, and manually measured stomatal-complex area from the same tomato cohort.",
        fill=LIGHT,
        accent=TEAL,
    )


def add_exec_summary(doc, area_stats, thermal_all, thermal_valid, plant_area):
    kicker(doc, "Decision brief")
    heading(doc, "Executive summary", 1)
    add_metric_strip(doc, [
        ("32", "PLANT UNITS", PALE_BLUE, BLUE),
        ("16", "VARIETIES", PALE_TEAL, TEAL),
        ("85", "MANUAL MASKS", PALE_GOLD, GOLD),
        ("14", "THERMAL QC PAIRS", LIGHT, NAVY),
    ])
    add_callout(
        doc,
        "Primary interpretation",
        "The current batch does not show a clear overall treatment effect on canopy temperature or stomatal-complex area. The strongest signal is a provisional increase in projected plant area under drought among thermal-mask-QC-valid pairs.",
        fill=PALE_TEAL,
    )
    heading(doc, "What the data show", 2)
    rich_paragraph(doc, [
        ("Canopy temperature. ", {"bold": True, "color": NAVY}),
        (
            f"Across all 16 pairs, drought-stressed plants averaged {thermal_all['drought_mean']:.2f} °C "
            f"versus {thermal_all['watered_mean']:.2f} °C for watered plants "
            f"(difference {thermal_all['difference']:+.2f} °C; p = {thermal_all['p']:.3f}). "
            "Two plant masks failed QC, so this full-set estimate is descriptive.",
            {},
        ),
    ])
    rich_paragraph(doc, [
        ("Projected plant area. ", {"bold": True, "color": NAVY}),
        (
            f"After restricting to {plant_area['n']} complete QC-valid pairs, drought-stressed plants had "
            f"{plant_area['pct']:.1f}% larger projected area on average "
            f"(difference {plant_area['difference']:,.0f} pixels²; 95% CI "
            f"{plant_area['ci'][0]:,.0f} to {plant_area['ci'][1]:,.0f}; p = {plant_area['p']:.003f}). "
            "This remains provisional pending camera-geometry review and physical calibration.",
            {},
        ),
    ])
    rich_paragraph(doc, [
        ("Stomatal-complex area. ", {"bold": True, "color": NAVY}),
        (
            f"Drought-stressed plants averaged 10,211.5 pixels² versus 10,968.5 pixels² for watered plants. "
            f"The paired difference was {area_stats['mean_paired_difference']:,.1f} pixels² "
            f"(-6.9%; 95% CI {area_stats['difference_95_ci_low']:,.1f} to "
            f"{area_stats['difference_95_ci_high']:,.1f}; p = {area_stats['paired_t_p_value']:.3f}).",
            {},
        ),
    ])
    rich_paragraph(doc, [
        ("Cross-scale relationship. ", {"bold": True, "color": NAVY}),
        (
            "Plant-level stomatal area was essentially unrelated to canopy temperature or projected plant "
            "area in this batch. The two pipelines are complementary rather than interchangeable.",
            {},
        ),
    ])
    heading(doc, "Recommended use", 2)
    add_key_value_table(doc, [
        ("Use now", "Descriptive treatment and variety profiles; candidate prioritization; protocol refinement."),
        ("Do not claim yet", "Generalized drought-response mechanisms, calibrated physical area, stomatal density, or variety ranking with inferential certainty."),
        ("Next milestone", "Correct two thermal masks, acquire physical scale references, and repeat with biological replication."),
    ])


def add_design(doc, thermal: pd.DataFrame):
    kicker(doc, "One cohort, two scales")
    heading(doc, "Study design and evidence chain", 1)
    paragraph(
        doc,
        "SeedThermal and StomatalPhenotypingPipeline are treated as a single experiment. "
        "Every row resolves to a stable plant unit, a variety, and one of two biological treatment labels.",
    )
    add_key_value_table(doc, [
        ("Watered", "Source code Regular; 100 mL."),
        ("Drought-stressed", "Source code Control; 0 mL."),
        ("Experimental unit", "One plant. Individual stomata are within-plant subsamples, not independent replicates."),
        ("Pairing", "One watered and one drought-stressed plant for each of 16 varieties."),
    ])
    heading(doc, "Evidence collected", 2)
    table = doc.add_table(rows=1, cols=4)
    headers = ["Scale", "Input", "Primary traits", "Usable scope"]
    for c, text in zip(table.rows[0].cells, headers):
        set_cell_shading(c, NAVY)
        set_run(c.paragraphs[0].add_run(text), 8.5, True, color=WHITE)
    rows = [
        ("Canopy", "64 radiometric FLIR captures across 13-15 July", "Mean, median, range, dispersion, ΔT-air", "July 15 isolates 32 plants; earlier paired frames show combined canopies."),
        ("Plant silhouette", "32 July 15 binary masks", "Projected area and frame fraction", "30 masks pass QC; 14 complete variety pairs."),
        ("Leaf epidermis", "32 micrographs", "Candidate confidence and per-object evidence", "5,633 candidates retained across decision tiers."),
        ("Stomatal complex", "85 manual polygons", "Area, length, width, perimeter, shape, orientation", "2-3 quality-selected objects per plant; pixel units."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for c, text in zip(cells, row):
            set_run(c.paragraphs[0].add_run(text), 8.2, color=INK)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            c.paragraphs[0].paragraph_format.line_spacing = 1.06
    set_table_geometry(table, [1300, 2300, 2600, 3160])
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    source_note(doc, "Source: combined plant manifest, thermal run table, calibrated candidate ledger, and manual mask manifest.")
    heading(doc, "Thermal acquisition context", 2)
    session = thermal.groupby("session_date").agg(
        captures=("mean_c", "size"), mean_c=("mean_c", "mean"), minimum=("mean_c", "min"), maximum=("mean_c", "max")
    ).reset_index()
    table = doc.add_table(rows=1, cols=5)
    for c, text in zip(table.rows[0].cells, ["Date", "Captures", "Mean °C", "Range °C", "Interpretation"]):
        set_cell_shading(c, TEAL)
        set_run(c.paragraphs[0].add_run(text), 8.5, True, color=WHITE)
    interpretations = [
        "Paired full-frame: both treatments in one image",
        "Paired full-frame: both treatments in one image",
        "Individual auto-masked plants; treatment comparison",
    ]
    for (_, row), interp in zip(session.iterrows(), interpretations):
        cells = table.add_row().cells
        vals = [
            str(row["session_date"]),
            f"{int(row['captures'])}",
            f"{row['mean_c']:.2f}",
            f"{row['minimum']:.2f}-{row['maximum']:.2f}",
            interp,
        ]
        for c, text in zip(cells, vals):
            set_run(c.paragraphs[0].add_run(text), 8.2)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
    set_table_geometry(table, [1500, 1000, 1100, 1500, 4260])
    set_table_borders(table)


def add_thermal(doc, thermal_all, thermal_valid):
    kicker(doc, "Canopy phenotype")
    heading(doc, "Thermal findings", 1)
    add_callout(
        doc,
        "Bottom line",
        f"No clear overall treatment-temperature difference was detected. Across all 16 pairs, the mean "
        f"drought-minus-watered difference was {thermal_all['difference']:+.2f} °C "
        f"(95% CI {thermal_all['ci'][0]:+.2f} to {thermal_all['ci'][1]:+.2f}; p = {thermal_all['p']:.3f}).",
        fill=PALE_BLUE,
        accent=BLUE,
    )
    caption(doc, 1, "Paired canopy temperatures by variety. Red QC labels mark the two pairs affected by failed automatic masks.")
    add_picture(doc, ASSETS / "thermal_variety_pairs.png", width=6.1)
    source_note(doc, "Source: July 15 plant-level FLIR extraction. Watered = Regular; drought-stressed = Control.")
    heading(doc, "QC-restricted sensitivity result", 2)
    paragraph(
        doc,
        f"After excluding NC264G and 26-033 to retain {thermal_valid['n']} complete QC-valid pairs, "
        f"drought-stressed plants averaged {thermal_valid['drought_mean']:.2f} °C and watered plants "
        f"{thermal_valid['watered_mean']:.2f} °C. The paired difference was "
        f"{thermal_valid['difference']:+.2f} °C (95% CI {thermal_valid['ci'][0]:+.2f} to "
        f"{thermal_valid['ci'][1]:+.2f}; paired t-test p = {thermal_valid['p']:.3f}; "
        f"Wilcoxon p = {thermal_valid['wilcoxon']:.3f}). This near-zero-to-modest cooling estimate "
        "should be treated as inconclusive rather than as evidence of equivalence.",
    )


def add_projected_area(doc, plant_area):
    kicker(doc, "Plant silhouette phenotype")
    heading(doc, "Projected plant area", 1)
    add_metric_strip(doc, [
        (f"{plant_area['watered_mean']:,.0f}", "WATERED MEAN PX²", PALE_BLUE, BLUE),
        (f"{plant_area['drought_mean']:,.0f}", "DROUGHT MEAN PX²", PALE_GOLD, GOLD),
        (f"+{plant_area['pct']:.1f}%", "PAIRED DIFFERENCE", PALE_TEAL, TEAL),
        (f"{plant_area['p']:.003f}", "PAIRED T-TEST P", LIGHT, NAVY),
    ])
    caption(doc, 2, "Within-variety difference in projected plant-mask area among the 14 complete thermal-QC pairs.")
    add_picture(doc, ASSETS / "projected_area_difference.png", width=6.1)
    source_note(doc, "Blue bars favor watered; ochre bars favor drought-stressed. Area is in image pixels, not physical units.")
    paragraph(
        doc,
        f"Drought-stressed plants were larger in {plant_area['positive_pairs']} of 14 QC-valid pairs. "
        f"The mean paired difference was {plant_area['difference']:,.0f} pixels² "
        f"(95% CI {plant_area['ci'][0]:,.0f} to {plant_area['ci'][1]:,.0f}); "
        f"paired t-test p = {plant_area['p']:.003f}, Wilcoxon p = {plant_area['wilcoxon']:.003f}.",
    )
    add_callout(
        doc,
        "Why this remains provisional",
        "The signal is statistically coherent within this batch, but projected pixel area depends on mask quality, camera distance, angle, and plant pose. It is not yet calibrated leaf area.",
        fill=PALE_GOLD,
        accent=GOLD,
    )


def add_stomata_method(doc, plant: pd.DataFrame):
    kicker(doc, "Object-level evidence")
    heading(doc, "Stomatal detection and measurement", 1)
    totals = plant[
        ["stomata_candidates_total", "stomata_candidates_accepted", "stomata_candidates_review", "stomata_candidates_rejected"]
    ].sum()
    add_metric_strip(doc, [
        (f"{int(totals['stomata_candidates_total']):,}", "CANDIDATES", LIGHT, NAVY),
        (f"{int(totals['stomata_candidates_accepted']):,}", "ACCEPTED TIER", PALE_TEAL, TEAL),
        (f"{int(totals['stomata_candidates_review']):,}", "REVIEW TIER", PALE_GOLD, GOLD),
        (f"{int(totals['stomata_candidates_rejected']):,}", "REJECTED TIER", PALE_BLUE, BLUE),
    ])
    paragraph(
        doc,
        "The automated system is structured as an analysis engine: it preserves every candidate, its model evidence, "
        "confidence-related fields, decision tier, image position, and measurement context. Candidate counts are a "
        "derived summary, not the primary scientific output.",
    )
    caption(doc, 3, "Example of a quality-selected stomatal complex with the manually traced outer contour.")
    add_picture(doc, ASSETS / "manual_mask_example.png", width=3.25)
    source_note(doc, "White boundary = reviewer polygon; translucent teal = measured stomatal-complex region.")
    heading(doc, "Manual measurement subset", 2)
    add_key_value_table(doc, [
        ("Coverage", "85 accepted manual polygons across 32 plants and all 16 varieties; 2-3 objects per plant."),
        ("Selection", "Highest local focus and contrast with low clipping and good edge containment among completed masks."),
        ("Not used for selection", "Treatment, variety, object area, mask area, or apparent stomatal shape."),
        ("Available traits", "Complex area, length, width, perimeter, aspect ratio, eccentricity, solidity, circularity, orientation, centroid."),
        ("Current unit", "Pixels and pixels². Physical dimensions require a stage-micrometer calibration."),
    ])


def add_stomatal_results(doc, area_stats):
    kicker(doc, "Microscopic phenotype")
    heading(doc, "Stomatal-complex area by treatment", 1)
    add_metric_strip(doc, [
        ("10,968.5", "WATERED PX²", PALE_BLUE, BLUE),
        ("10,211.5", "DROUGHT PX²", PALE_GOLD, GOLD),
        ("-6.9%", "RELATIVE DIFFERENCE", PALE_TEAL, TEAL),
        (f"{area_stats['paired_t_p_value']:.3f}", "PAIRED T-TEST P", LIGHT, NAVY),
    ])
    caption(doc, 4, "Plant-level mean stomatal-complex area for each of 16 matched variety pairs.")
    add_picture(doc, ASSETS / "stomatal_variety_pairs.png", width=6.1)
    source_note(doc, "Each plant mean is based on 2-3 quality-selected manual masks; inferential tests use plants, not individual stomata.")
    paragraph(
        doc,
        f"The mean paired difference, defined as drought-stressed minus watered, was "
        f"{area_stats['mean_paired_difference']:,.1f} pixels². Its 95% confidence interval ranged from "
        f"{area_stats['difference_95_ci_low']:,.1f} to {area_stats['difference_95_ci_high']:,.1f} pixels². "
        f"The paired t-test gave p = {area_stats['paired_t_p_value']:.3f}; the Wilcoxon signed-rank test "
        f"gave p = {area_stats['wilcoxon_p_value']:.3f}; Cohen's dz = {area_stats['cohen_dz']:.2f}. "
        "The estimate is compatible with both modest drought-associated reduction and modest increase.",
    )
    add_callout(
        doc,
        "Interpretation",
        "There is no clear batch-level treatment effect on stomatal-complex area. The variety-specific directions are heterogeneous and should guide replication, not selection decisions.",
        fill=PALE_TEAL,
    )


def add_variety_results(doc, variety: pd.DataFrame):
    kicker(doc, "Genotype profiles")
    heading(doc, "Descriptive comparison among 16 varieties", 1)
    paragraph(
        doc,
        "The highest two-treatment descriptive mean was observed for 26-037; the lowest for NC264G. "
        "Because each variety-treatment cell contains one plant, these are profiles of this cohort, not population estimates.",
    )
    table = doc.add_table(rows=1, cols=5)
    headers = ["Variety", "Watered px²", "Drought px²", "Two-treatment mean", "Drought - watered"]
    for c, text in zip(table.rows[0].cells, headers):
        set_cell_shading(c, NAVY)
        set_run(c.paragraphs[0].add_run(text), 7.7, True, color=WHITE)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    ranked = variety.sort_values("descriptive_mean_area", ascending=False)
    for i, (_, row) in enumerate(ranked.iterrows()):
        cells = table.add_row().cells
        if i % 2:
            for c in cells:
                set_cell_shading(c, LIGHT)
        values = [
            row["variety"],
            f"{row['watered_mean_area']:,.0f}",
            f"{row['drought_mean_area']:,.0f}",
            f"{row['descriptive_mean_area']:,.0f}",
            f"{row['difference_drought_minus_watered']:+,.0f}",
        ]
        for j, (c, text) in enumerate(zip(cells, values)):
            set_run(c.paragraphs[0].add_run(str(text)), 7.7, bold=(j == 0), color=INK)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.RIGHT
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            c.paragraphs[0].paragraph_format.line_spacing = 1.0
            set_cell_margins(c, top=55, bottom=55, start=95, end=95)
    set_table_geometry(table, [2500, 1600, 1600, 1900, 1760])
    set_table_borders(table, size=4)
    set_repeat_table_header(table.rows[0])
    source_note(doc, "Ranking is based on the descriptive average of the watered and drought-stressed plant means.")
    top_inc = variety.loc[variety["difference_drought_minus_watered"].idxmax()]
    top_dec = variety.loc[variety["difference_drought_minus_watered"].idxmin()]
    add_callout(
        doc,
        "Response range",
        f"The largest positive response was {top_inc['variety']} ({top_inc['difference_drought_minus_watered']:+,.0f} pixels²); "
        f"the largest negative response was {top_dec['variety']} ({top_dec['difference_drought_minus_watered']:+,.0f} pixels²). "
        "These extremes are especially sensitive to low replication.",
        fill=PALE_GOLD,
        accent=GOLD,
    )


def add_integrated(doc, cross, valid):
    kicker(doc, "Cross-scale synthesis")
    heading(doc, "Do canopy and stomatal traits move together?", 1)
    caption(doc, 5, "Plant-level stomatal-complex area versus thermal and projected-area traits for the 30 thermal-QC-valid plants.")
    add_picture(doc, ASSETS / "cross_modal_scatter.png", width=6.15)
    source_note(doc, "Correlations are exploratory; points are plants and colors indicate treatment.")
    add_key_value_table(doc, [
        ("Stomatal area vs canopy temperature",
         f"Pearson r = {cross['temp'].statistic:.03f}; p = {cross['temp'].pvalue:.3f}; n = {cross['n']} plants."),
        ("Stomatal area vs projected plant area",
         f"Pearson r = {cross['projected'].statistic:.03f}; p = {cross['projected'].pvalue:.3f}; n = {cross['n']} plants."),
        ("Paired stomatal response vs thermal response",
         f"Pearson r = {cross['response_temp'].statistic:.03f}; p = {cross['response_temp'].pvalue:.3f}; n = {len(valid)} varieties."),
        ("Paired stomatal response vs projected-area response",
         f"Pearson r = {cross['response_area'].statistic:.03f}; p = {cross['response_area'].pvalue:.3f}; n = {len(valid)} varieties."),
    ])
    paragraph(
        doc,
        "The cross-sectional relationships are essentially null. The response-level associations are moderate in "
        "direction but imprecise, especially for projected area (r = 0.47, p = 0.089). They are useful hypotheses "
        "for a replicated experiment, not evidence of mechanistic coupling.",
    )
    add_callout(
        doc,
        "Scientific reading",
        "A plant can show canopy-scale size or temperature differences without a matching change in the area of its clearest stomatal complexes. Response is multiscale; one trait cannot substitute for another.",
        fill=PALE_TEAL,
    )


def add_limits_and_next(doc):
    kicker(doc, "Interpretation guardrails")
    heading(doc, "Limitations and recommended next steps", 1)
    table = doc.add_table(rows=1, cols=3)
    for c, text in zip(table.rows[0].cells, ["Constraint", "Why it matters", "Recommended action"]):
        set_cell_shading(c, NAVY)
        set_run(c.paragraphs[0].add_run(text), 8.2, True, color=WHITE)
    rows = [
        ("One plant per variety-treatment cell", "Variety effects and interactions cannot be estimated reliably.", "Repeat with independent biological replicates."),
        ("2-3 selected stomata per plant", "Within-leaf variation is only sparsely represented.", "Expand objects per plant using the validated analysis engine."),
        ("Quality-selected objects", "Results describe clearest measurable stomata, not a random within-image population.", "Predefine random/systematic fields and selection rules."),
        ("Pixel-only geometry", "Areas cannot be compared in µm² or across changed magnification.", "Acquire a stage-micrometer image at every optical configuration."),
        ("Two failed thermal masks", "NC264G watered and 26-033 drought distort area and temperature extraction.", "Correct masks and rerun the joined analysis."),
        ("Projected area is view-dependent", "Camera pose and plant architecture affect silhouette size.", "Fix camera geometry and add a scale reference."),
        ("No stomatal density or pore aperture", "Current masks support complex morphology, not density or conductance.", "Define analyzed leaf area; validate pore segmentation separately."),
    ]
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        if i % 2:
            for c in cells:
                set_cell_shading(c, LIGHT)
        for c, text in zip(cells, row):
            set_run(c.paragraphs[0].add_run(text), 8.1)
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            c.paragraphs[0].paragraph_format.line_spacing = 1.06
    set_table_geometry(table, [2300, 3300, 3760])
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    source_note(doc, "These constraints define the boundary between current evidence and future claims.")
    heading(doc, "Prioritized next experiment", 2)
    add_key_value_table(doc, [
        ("1  Repair", "Correct the two thermal masks and regenerate all joined outputs."),
        ("2  Calibrate", "Capture stage-micrometer and plant-view scale references."),
        ("3  Replicate", "Target at least 3-5 plants per variety-treatment cell; preserve within-variety pairing."),
        ("4  Sample", "Use predefined fields and a consistent number of stomata per plant."),
        ("5  Validate", "Reserve held-out tomato images for detection calibration and mask-error assessment."),
        ("6  Extend", "Only after validation, add pore aperture, density, and spatial-pattern traits."),
    ])
    add_callout(
        doc,
        "Go/no-go recommendation",
        "Proceed with the combined platform as a research workflow. Do not yet use this single batch to rank varieties for drought tolerance or claim a stomatal mechanism.",
        fill=PALE_GOLD,
        accent=GOLD,
    )


def add_methods_and_provenance(doc):
    kicker(doc, "Reproducibility")
    heading(doc, "Methods and data provenance", 1)
    heading(doc, "Thermal pipeline", 2)
    paragraph(
        doc,
        "Radiometric FLIR JPGs were converted to per-pixel Celsius arrays. July 15 plant masks defined the "
        "canopy pixels used for mean, median, minimum, maximum, standard deviation, P10, P90, and temperature "
        "relative to the 25 °C ambient reference. Projected plant area is the number and fraction of selected pixels.",
    )
    heading(doc, "Stomatal pipeline", 2)
    paragraph(
        doc,
        "Native-resolution micrographs were processed with a high-recall, tiled candidate detector. The object "
        "ledger retains decision tiers and confidence evidence. For this report, manual outer stomatal-complex "
        "polygons were ranked within each image by local focus, contrast, clipping, and edge containment; at most "
        "three objects per plant were retained. Polygon geometry produced direct pixel-space morphology.",
    )
    heading(doc, "Statistical analysis", 2)
    paragraph(
        doc,
        "Treatment contrasts were paired by variety. Plant means, not individual stomata, were the inferential "
        "units. Paired t-tests were accompanied by 95% confidence intervals and Wilcoxon signed-rank sensitivity "
        "tests. Variety results are descriptive because each variety-treatment cell contains one plant. "
        "Cross-modal Pearson correlations were exploratory and restricted to thermal-mask-QC-valid plants or pairs.",
    )
    heading(doc, "Primary source files", 2)
    files = [
        ("Combined plant table", "results/combined_quality_mask_analysis_v1/plant_phenotypes.csv"),
        ("Variety-level paired table", "results/combined_quality_mask_analysis_v1/variety_phenotypes.csv"),
        ("Per-stoma measurements", "results/quality_mask_analysis_v1/stomatal_complex_measurements.csv"),
        ("Stomatal statistical record", "results/quality_mask_analysis_v1/area_comparisons/stomatal_area_analysis.json"),
        ("Thermal run table", "../SeedThermal/processed/thermal_data_all_sessions.csv"),
        ("Study definitions", "docs/COMBINED_STUDY.md and docs/ANALYSIS_ENGINE.md"),
    ]
    add_key_value_table(doc, files)
    heading(doc, "Reporting language", 2)
    add_callout(
        doc,
        "Preferred conclusion",
        "In this quality-selected pilot cohort, drought stress was not associated with a clear overall change in canopy temperature or stomatal-complex area. A larger projected plant silhouette was observed in the QC-valid drought group, but this finding remains provisional pending mask correction, calibration, and biological replication.",
        fill=PALE_TEAL,
    )


def build_report():
    plant = pd.read_csv(PLANT_CSV)
    variety = pd.read_csv(VARIETY_CSV)
    stomata = pd.read_csv(STOMATA_CSV)
    thermal = pd.read_csv(THERMAL_CSV)
    generate_figures(plant, variety, stomata)
    area_stats, thermal_all, thermal_valid, plant_area, cross, valid = compute_results(plant, variety)

    doc = Document()
    style_document(doc)
    add_cover(doc)
    add_section_break(doc)
    add_exec_summary(doc, area_stats, thermal_all, thermal_valid, plant_area)
    add_section_break(doc)
    add_design(doc, thermal)
    add_section_break(doc)
    add_thermal(doc, thermal_all, thermal_valid)
    add_section_break(doc)
    add_projected_area(doc, plant_area)
    add_section_break(doc)
    add_stomata_method(doc, plant)
    add_section_break(doc)
    add_stomatal_results(doc, area_stats)
    add_section_break(doc)
    add_variety_results(doc, variety)
    add_section_break(doc)
    add_integrated(doc, cross, valid)
    add_section_break(doc)
    add_limits_and_next(doc)
    add_section_break(doc)
    add_methods_and_provenance(doc)

    props = doc.core_properties
    props.title = "From Canopy Heat to Stomatal Form"
    props.subject = "Integrated SeedThermal and StomatalPhenotypingPipeline report"
    props.author = "Sapkota Lab"
    props.keywords = "tomato, drought, thermal phenotyping, stomata, plant area"
    props.comments = "Generated from versioned project outputs on 29 July 2026."
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build_report()
